#!/usr/bin/env python3
"""
conformance.py - Tests internes de conformite a la charte The Oz V2.

Verifie un .pptx genere par le moteur, regle par regle, en inspectant reellement
le fichier (formes, couleurs, polices, alignements) - pas le code qui l'a produit.
Prend en compte l'heritage du masque (bande, trait de titre, pied).

Usage :
    python3 tests/conformance.py out/showcase.pptx [autre.pptx ...]
    python3 tests/conformance.py --all      # rend puis controle tous les plans

Sortie : un rapport lisible ; code de sortie 1 si une regle est violee.
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "engine"))

from pptx import Presentation                                    # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE                      # noqa: E402
from pptx.enum.text import PP_ALIGN                              # noqa: E402
from oz_deck.theme import Brand, CHARTE_VERSION                  # noqa: E402
from oz_deck import docprops, __version__ as ENGINE_VERSION      # noqa: E402
from oz_deck.theme import export_tokens                          # noqa: E402

EMU = 914400.0
TOL = 0.02   # tolerance geometrique (pouces)

HEX = lambda c: f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"  # noqa: E731

PALETTE = {
    "191919": "noir", "FFFFFF": "blanc", "EC4324": "orange logo",
    "3E9356": "vert sauge", "D3122A": "rubis", "999999": "gris",
    "D99B26": "moutarde", "EA7C23": "citrouille", "8F3D2E": "cuivre",
    "59261D": "automne", "F3F2F2": "gris technique", "E7E7E7": "filet",
}
CAMAIEU = {"D99B26", "EA7C23", "EC4324", "8F3D2E", "59261D"}
EVENT_BG = "EC4324"
BANNED_PUNCT = ("—", "–")   # cadratin, demi-cadratin

# Formes ou le centrage est admis (titres de colonne, badges, pastilles, onglets).
CENTER_OK_PREFIX = ("badge", "pastille", "onglet", "phase", "visuel")


class Report:
    def __init__(self, path: Path):
        self.path = path
        self.rows: list[tuple[str, bool, str]] = []

    def check(self, rule: str, ok: bool, detail: str = "") -> bool:
        self.rows.append((rule, ok, detail))
        return ok

    @property
    def failed(self) -> list[tuple[str, bool, str]]:
        return [r for r in self.rows if not r[1]]

    def render(self) -> str:
        out = [f"\n=== {self.path.name} ==="]
        for rule, ok, detail in self.rows:
            mark = "OK  " if ok else "KO  "
            out.append(f"  {mark}{rule}" + (f"  -> {detail}" if detail else ""))
        return "\n".join(out)


# ----------------------------- introspection -----------------------------

def _inch(v) -> float:
    return (v or 0) / EMU


def _shape_fill_hex(sh) -> str | None:
    try:
        if sh.fill.type is not None and sh.fill.type == 1:      # MSO_FILL.SOLID
            return HEX(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def _iter_shapes(container):
    for sh in container.shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in _iter_shapes(sh):
                yield sub


def _effective_shapes(slide):
    """Formes reellement visibles : masque (si non masque) + disposition + slide."""
    layout = slide.slide_layout
    show_master = (layout.element.get("showMasterSp") or "1") != "0"
    shapes = []
    if show_master:
        shapes += list(_iter_shapes(layout.slide_master))
    shapes += list(_iter_shapes(layout))
    shapes += list(_iter_shapes(slide))
    return shapes


def _bg_hex(part) -> str | None:
    """Couleur de fond declaree sur une slide / disposition / masque."""
    el = part.element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld")
    if el is None:
        return None
    bg = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if bg is None:
        return None
    clr = bg.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    return clr.get("val").upper() if clr is not None else None


def _effective_bg(slide) -> str | None:
    layout = slide.slide_layout
    for part in (slide, layout, layout.slide_master):
        hexv = _bg_hex(part)
        if hexv:
            return hexv
    return None


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _title_ph_el(part):
    for sp in part.element.iter(f"{_P}sp"):
        ph = sp.find(f".//{_P}nvSpPr/{_P}nvPr/{_P}ph")
        if ph is not None and ph.get("type") in ("title", "ctrTitle"):
            return sp
    return None


def _effective_caps(slide, title_shape) -> str | None:
    """Resout la capitalisation reellement appliquee au titre : rPr du run, puis
    lstStyle de la slide, de la disposition, puis titleStyle du masque."""
    for r in title_shape.text_frame.paragraphs[0].runs:
        rpr = r._r.find(f"{_A}rPr")
        if rpr is not None and rpr.get("cap"):
            return rpr.get("cap")
    layout = slide.slide_layout
    for part in (slide, layout):
        sp = _title_ph_el(part)
        if sp is None:
            continue
        el = sp.find(f".//{_A}lstStyle/{_A}lvl1pPr/{_A}defRPr")
        if el is not None and el.get("cap"):
            return el.get("cap")
    el = layout.slide_master.element.find(
        f".//{_P}txStyles/{_P}titleStyle/{_A}lvl1pPr/{_A}defRPr")
    return el.get("cap") if el is not None else None


_FONTS: dict = {}


def _pil_font(size_pt: float, bold: bool):
    """Charge la vraie Quicksand pour mesurer le texte (a defaut : None)."""
    key = (round(size_pt, 1), bold)
    if key in _FONTS:
        return _FONTS[key]
    try:
        from PIL import ImageFont
        name = "Quicksand-Bold.ttf" if bold else "Quicksand-Regular.ttf"
        path = ROOT / "engine" / "oz_deck" / "assets" / "fonts" / name
        # 96 dpi : 1 pt = 96/72 px
        f = ImageFont.truetype(str(path), int(round(size_pt * 96 / 72)))
    except Exception:
        f = None
    _FONTS[key] = f
    return f


def _text_height_in(shape) -> float | None:
    """Hauteur rendue estimee du texte d'une forme, en pouces (None si non mesurable).

    Mesure avec les fontes Quicksand du depot : retour a la ligne sur la largeur
    utile, puis somme des lignes a l'interlignage du paragraphe.
    """
    try:
        tf = shape.text_frame
    except Exception:
        return None
    if not tf.text.strip():
        return None
    width_in = _inch(shape.width) - _inch(tf.margin_left) - _inch(tf.margin_right)
    if width_in <= 0.05:
        return None
    total = 0.0
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = max((r.font.size.pt for r in runs if r.font.size), default=None)
        if size is None:
            return None
        bold = any(r.font.bold for r in runs)
        font = _pil_font(size, bold)
        if font is None:
            return None
        text = "".join(r.text for r in runs)
        if not tf.word_wrap:
            lines = 1
        else:
            limit_px = width_in * 96
            lines, cur = 1, ""
            for word in text.split(" "):
                trial = (cur + " " + word).strip()
                if font.getlength(trial) <= limit_px or not cur:
                    cur = trial
                else:
                    lines += 1
                    cur = word
        ls = float(p.line_spacing) if p.line_spacing else 1.2
        total += lines * size * ls / 72.0
        if p.space_after is not None:
            total += p.space_after.pt / 72.0
    return total


def _has_theme_style(sh) -> bool:
    """Vrai si la forme porte un `<p:style>` : elle herite alors des effets du
    theme Office, dont l'ombre portee que la charte interdit."""
    try:
        return sh._element.find(f"{_P}style") is not None
    except Exception:
        return False


def _effective_title_size(slide, title_shape) -> float | None:
    """Corps reellement applique au titre : run, puis lstStyle de la slide et de la
    disposition, puis titleStyle du masque."""
    for r in title_shape.text_frame.paragraphs[0].runs:
        if r.font.size:
            return r.font.size.pt
    layout = slide.slide_layout
    for part in (slide, layout):
        sp = _title_ph_el(part)
        if sp is None:
            continue
        el = sp.find(f".//{_A}lstStyle/{_A}lvl1pPr/{_A}defRPr")
        if el is not None and el.get("sz"):
            return int(el.get("sz")) / 100.0
    el = layout.slide_master.element.find(
        f".//{_P}txStyles/{_P}titleStyle/{_A}lvl1pPr/{_A}defRPr")
    return int(el.get("sz")) / 100.0 if (el is not None and el.get("sz")) else None


def _runs(slide):
    """(shape, paragraph, run) de toutes les formes de la slide."""
    for sh in _iter_shapes(slide):
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                yield sh, p, r


def _full_bleed(sh, hexv=None) -> bool:
    try:
        return (_inch(sh.width) > Brand.SW - 0.2 and _inch(sh.height) > Brand.SH - 0.2
                and (hexv is None or _shape_fill_hex(sh) == hexv))
    except Exception:
        return False


def _event_backdrop(slide):
    """Fond de slide evenementielle : image de degrade en pleine page, aplat
    orange plein ecran, ou fond de page orange. Retourne la forme, True, ou None."""
    for sh in _iter_shapes(slide):
        if not _full_bleed(sh):
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return sh
        if _shape_fill_hex(sh) == EVENT_BG:
            return sh
    return True if _effective_bg(slide) == EVENT_BG else None


def _is_event(slide) -> bool:
    return _event_backdrop(slide) is not None


# ----------------------------- regles -----------------------------

def check_deck(path: Path) -> Report:
    rep = Report(path)
    prs = Presentation(str(path))

    # 1. format 16:9 paysage
    w, h = _inch(prs.slide_width), _inch(prs.slide_height)
    rep.check("format 16:9 paysage obligatoire",
              abs(w - Brand.SW) < TOL and abs(h - Brand.SH) < TOL,
              f"{w:.3f} x {h:.3f} pouces")

    bad_palette, bad_font, bad_align, bad_line, bad_punct = [], [], [], [], []
    bad_spacing, bad_caps = [], []
    missing_band, missing_rule, missing_footer, bad_bg, missing_page = [], [], [], [], []
    bad_chart, bad_logo, overflow = [], [], []
    img_out, img_squashed = [], []
    shadows, too_small, bad_title_size = [], [], []
    n_event = n_content = 0

    for i, slide in enumerate(prs.slides, start=1):
        event = _is_event(slide)
        n_event += event
        n_content += (not event)
        eff = _effective_shapes(slide)
        tag = f"slide {i}"

        # --- fonds ---
        if event:
            if _event_backdrop(slide) is None:
                bad_bg.append(f"{tag}: evenementielle sans fond de page")
        else:
            bg = _effective_bg(slide)
            if bg != "FFFFFF":
                bad_bg.append(f"{tag}: fond de contenu {bg} (attendu FFFFFF)")
            # bande orange 0,5 cm au bord gauche
            band = [sh for sh in eff
                    if _shape_fill_hex(sh) == "EC4324"
                    and _inch(sh.left) < TOL
                    and abs(_inch(sh.width) - Brand.BAND_W) < 0.03
                    and _inch(sh.height) > Brand.SH - 0.2]
            if not band:
                missing_band.append(tag)
            # trait orange 3 pt sous le titre
            rule = [sh for sh in eff
                    if _shape_fill_hex(sh) == "EC4324"
                    and abs(_inch(sh.top) - Brand.RULE_Y) < 0.06
                    and abs(_inch(sh.width) - Brand.RULE_W) < 0.06
                    and abs(_inch(sh.height) - Brand.RULE_PT / 72.0) < 0.02]
            if not rule:
                missing_rule.append(tag)

        # --- pied de page ---
        texts = [r.text.strip().upper() for _, _, r in _runs(slide)]
        eff_texts = []
        for sh in eff:
            if sh.has_text_frame:
                eff_texts.append(sh.text_frame.text.strip().upper())
        if not any(t.startswith("THE OZ - AGENCE SHOPIFY") for t in eff_texts):
            missing_footer.append(f"{tag}: mention gauche absente")
        filet = [sh for sh in eff
                 if abs(_inch(sh.top) - Brand.FOOTER_LINE_Y) < 0.03
                 and _inch(sh.width) > Brand.SW - 1.2
                 and _inch(sh.height) < 0.03]
        if not filet:
            missing_footer.append(f"{tag}: filet de pied absent")
        # La pagination est desormais posee sur TOUTES les slides.
        paginated = any(t.startswith(Brand.FOOTER_RIGHT + " - ") for t in texts)
        if not paginated:
            missing_page.append(tag)

        # --- titre en capitales (texte deja capitalise OU style cap="all" herite) ---
        title = slide.shapes.title
        if (not event) and title is not None and title.text_frame.text.strip():
            t = title.text_frame.text
            if t != t.upper() and _effective_caps(slide, title) != "all":
                bad_caps.append(f"{tag}: {t[:40]!r}")

        # --- logo jamais noir ---
        for sh in eff:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                nm = (sh.image.filename or sh.name or "").lower()
                if "noir" in nm or "black" in nm:
                    bad_logo.append(f"{tag}: {nm}")

        # --- aucune ombre portee (retour de relecture) ---
        for sh in _iter_shapes(slide):
            if not _has_theme_style(sh):
                continue          # textboxes et images : aucun style de theme a heriter
            try:
                if sh.shadow.inherit:
                    shadows.append(f"{tag}: {sh.name}")
            except Exception:
                pass

        # --- corps du titre de slide ---
        if not event:
            t = slide.shapes.title
            size = _effective_title_size(slide, t) if t is not None else None
            if size is not None and abs(size - Brand.TITLE) > 0.01:
                bad_title_size.append(f"{tag}: corps {size}")

        # --- images : confinees dans la slide, ratio d'aspect preserve ---
        for sh in _iter_shapes(slide):
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            l, t = _inch(sh.left), _inch(sh.top)
            w, h = _inch(sh.width), _inch(sh.height)
            limit = Brand.SH if event else Brand.FOOTER_LINE_Y
            if t + h > limit + 0.02 or l + w > Brand.SW + 0.02 or l < -0.02 or t < -0.02:
                img_out.append(f"{tag}: image {w:.2f}x{h:.2f}\" en ({l:.2f},{t:.2f}) "
                               f"depasse (limite basse {limit}\")")
            try:
                pw, ph = sh.image.size
                native = pw / ph
                drawn = w / h if h else native
                if abs(drawn - native) / native > 0.02:
                    img_squashed.append(f"{tag}: ratio {drawn:.3f} au lieu de "
                                        f"{native:.3f} ({sh.name})")
            except Exception:
                pass

        # --- debordement : le texte ne doit pas franchir le filet de pied ---
        for sh in _iter_shapes(slide):
            if not sh.has_text_frame:
                continue
            if _inch(sh.top) >= Brand.FOOTER_LINE_Y - 0.05:
                continue            # le pied vit sous le filet, par construction
            h = _text_height_in(sh)
            if h is None:
                continue
            bottom = _inch(sh.top) + h
            if bottom > Brand.FOOTER_LINE_Y - 0.04:
                overflow.append(f"{tag}: {sh.text_frame.text[:26]!r} descend a "
                                f"{bottom:.2f}\" (pied a {Brand.FOOTER_LINE_Y})")

        # --- couleurs, police, ferrage, interlignage, ponctuation ---
        for sh in _iter_shapes(slide):
            fh = _shape_fill_hex(sh)
            if fh and fh not in PALETTE:
                bad_palette.append(f"{tag}: aplat #{fh} ({sh.name})")
            try:
                if sh.line.fill.type == 1:
                    lh = HEX(sh.line.color.rgb)
                    if lh not in PALETTE:
                        bad_palette.append(f"{tag}: contour #{lh} ({sh.name})")
            except Exception:
                pass

        for sh, p, r in _runs(slide):
            if r.font.name and r.font.name != Brand.FONT:
                bad_font.append(f"{tag}: police {r.font.name!r}")
            try:
                ch = HEX(r.font.color.rgb)
                if ch not in PALETTE:
                    bad_palette.append(f"{tag}: texte #{ch} ({r.text[:22]!r})")
            except Exception:
                pass
            # plancher de corps : aucun texte sous 12 (le pied reste a 8, charte)
            size_pt = r.font.size.pt if r.font.size else None
            if (size_pt is not None and size_pt < Brand.MIN_TEXT - 0.01
                    and abs(size_pt - Brand.FOOTER) > 0.01):
                too_small.append(f"{tag}: corps {size_pt} sur {r.text[:22]!r}")
            if any(b in r.text for b in BANNED_PUNCT):
                bad_punct.append(f"{tag}: {r.text[:40]!r}")
            if p.alignment == PP_ALIGN.JUSTIFY:
                bad_align.append(f"{tag}: texte justifie ({r.text[:30]!r})")
            if p.alignment == PP_ALIGN.CENTER:
                nm = (sh.name or "").lower()
                short = len(r.text) <= 24
                allowed = short or any(k in nm for k in CENTER_OK_PREFIX)
                if not allowed:
                    bad_align.append(f"{tag}: bloc centre ({r.text[:30]!r})")
            # texte courant : interlignage strictement entre 1,3 et 1,4
            size = r.font.size.pt if r.font.size else None
            if size and 12 <= size <= 15 and not r.font.bold:
                ls = p.line_spacing
                if ls is None or not (1.3 - 1e-6 <= float(ls) <= 1.4 + 1e-6):
                    bad_line.append(f"{tag}: interlignage {ls} sur {r.text[:26]!r}")
            # espacement : avant = 2 x apres
            if p.space_before is not None and p.space_after is not None:
                b, a = p.space_before.pt, p.space_after.pt
                if a > 0 and abs(b - 2 * a) > 0.01:
                    bad_spacing.append(f"{tag}: {b}pt avant / {a}pt apres")

        # --- graphiques natifs, series au camaieu ---
        for sh in _iter_shapes(slide):
            if not getattr(sh, "has_chart", False):
                continue
            for si, ser in enumerate(sh.chart.series):
                for hexv in _series_colors(ser):
                    if hexv not in CAMAIEU | {"FFFFFF"}:
                        bad_chart.append(f"{tag}: serie {si} #{hexv}")

    rep.check("fonds : contenu blanc / evenementiel orange uni", not bad_bg,
              "; ".join(bad_bg[:3]))
    rep.check("bande orange 0,5 cm au bord gauche (slides de contenu)",
              not missing_band, "; ".join(missing_band[:5]))
    rep.check("trait orange 3 pt sous le titre", not missing_rule,
              "; ".join(missing_rule[:5]))
    rep.check("titres de contenu en capitales", not bad_caps, "; ".join(bad_caps[:3]))
    rep.check("pied standardise (filet 1 pt + logo + mention corps 8)",
              not missing_footer, "; ".join(missing_footer[:3]))
    rep.check("pagination sur toutes les slides", not missing_page,
              "; ".join(missing_page[:5]))
    rep.check("palette : aucune couleur hors charte", not bad_palette,
              "; ".join(sorted(set(bad_palette))[:4]))
    rep.check(f"typographie : {Brand.FONT} partout (aucune serif)", not bad_font,
              "; ".join(sorted(set(bad_font))[:3]))
    rep.check("ferrage a gauche : ni justifie ni bloc centre", not bad_align,
              "; ".join(bad_align[:3]))
    rep.check("interlignage du texte courant entre 1,3 et 1,4", not bad_line,
              "; ".join(bad_line[:3]))
    rep.check("espacement des titres : avant = 2 x apres", not bad_spacing,
              "; ".join(sorted(set(bad_spacing))[:3]))
    rep.check("graphiques natifs au camaieu utilitaire", not bad_chart,
              "; ".join(sorted(set(bad_chart))[:3]))
    rep.check("logo jamais noir", not bad_logo, "; ".join(sorted(set(bad_logo))[:3]))
    rep.check("aucun texte ne franchit le filet de pied", not overflow,
              "; ".join(overflow[:4]))
    rep.check("images contenues dans la zone utile", not img_out,
              "; ".join(img_out[:3]))
    rep.check("images non deformees (ratio d'aspect preserve)", not img_squashed,
              "; ".join(img_squashed[:3]))
    rep.check("aucune ombre portee sur les formes", not shadows,
              "; ".join(sorted(set(shadows))[:4]))
    rep.check(f"titres de slide au corps {Brand.TITLE:g}", not bad_title_size,
              "; ".join(bad_title_size[:3]))
    rep.check(f"aucun texte sous le corps {Brand.MIN_TEXT:g} (hors pied a "
              f"{Brand.FOOTER:g})", not too_small, "; ".join(sorted(set(too_small))[:4]))
    rep.check("ponctuation : pas de cadratin ni demi-cadratin", not bad_punct,
              "; ".join(bad_punct[:3]))
    rep.check("police Quicksand embarquee dans le fichier",
              _fonts_embedded(path), "")
    # Tampon de tracabilite : sans ce controle, il sauterait au premier refactor
    # du builder sans que personne ne le voie.
    stamp = docprops.read(str(path))
    want = {"OzCharteVersion": CHARTE_VERSION, "OzEngineVersion": ENGINE_VERSION}
    rep.check(f"version de charte tamponnee ({CHARTE_VERSION})",
              stamp == want,
              f"lu {stamp or 'aucune propriete personnalisee'}, attendu {want}")
    defects = _ooxml_defects(path)
    rep.check("structure OOXML valide (PowerPoint n'a rien a reparer)",
              not defects, "; ".join(defects[:3]))
    rep.check(f"slides : {n_content} de contenu, {n_event} evenementielles", True, "")
    return rep


def _series_colors(ser) -> list[str]:
    out = []
    for attr in ("fill", "line"):
        try:
            fmt = getattr(ser.format, attr)
            col = fmt.fore_color if attr == "fill" else fmt.color
            if col.type is not None:
                out.append(HEX(col.rgb))
        except Exception:
            pass
    return out


# --- validite structurelle OOXML ------------------------------------------
#
# ECMA-376 impose une SEQUENCE stricte aux enfants de ces elements : un ordre
# invalide donne un fichier que PowerPoint propose de "reparer", alors que le zip
# et le XML restent parfaitement bien formes. C'est exactement le piege quand on
# fabrique du XML a la main (masque, theme, fontembed).
_ORDERS = {
    # CT_TextParagraphProperties : pPr, lvl1pPr..lvl9pPr, defPPr
    "pPr": ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct",
            "buSzPts", "buFontTx", "buFont", "buNone", "buAutoNum", "buChar",
            "tabLst", "defRPr", "extLst"],
    # CT_TextCharacterProperties : rPr, defRPr, endParaRPr
    "rPr": ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill",
            "grpFill", "effectLst", "effectDag", "highlight", "uLnTx", "uLn",
            "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick",
            "hlinkMouseOver", "rtl", "extLst"],
    # CT_ShapeProperties : spPr (a:, p: et c:)
    "spPr": ["xfrm", "custGeom", "prstGeom", "noFill", "solidFill", "gradFill",
             "blipFill", "pattFill", "grpFill", "ln", "effectLst", "effectDag",
             "scene3d", "sp3d", "extLst"],
    # CT_TextBody : txBody, txPr
    "txBody": ["bodyPr", "lstStyle", "p"],
    "sp": ["nvSpPr", "spPr", "style", "txBody"],
    "pic": ["nvPicPr", "blipFill", "spPr", "style"],
    "nvSpPr": ["cNvPr", "cNvSpPr", "nvPr"],
    "nvPicPr": ["cNvPr", "cNvPicPr", "nvPr"],
    "bg": ["bgPr", "bgRef"],
    "bgPr": ["noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
             "effectLst", "effectDag"],
}
_ORDERS["lvl1pPr"] = _ORDERS["lvl2pPr"] = _ORDERS["lvl3pPr"] = _ORDERS["pPr"]
_ORDERS["lvl4pPr"] = _ORDERS["lvl5pPr"] = _ORDERS["lvl6pPr"] = _ORDERS["pPr"]
_ORDERS["lvl7pPr"] = _ORDERS["lvl8pPr"] = _ORDERS["lvl9pPr"] = _ORDERS["pPr"]
_ORDERS["defPPr"] = _ORDERS["pPr"]
_ORDERS["defRPr"] = _ORDERS["endParaRPr"] = _ORDERS["rPr"]
_ORDERS["txPr"] = _ORDERS["txBody"]

# Attributs obligatoires souvent oublies quand on ecrit l'element a la main.
_REQUIRED_ATTRS = {"numFmt": ["formatCode"]}


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _ooxml_defects(path: Path) -> list[str]:
    """Verifie l'ordre des enfants et les attributs requis dans toutes les parts."""
    import zipfile
    from lxml import etree
    out = []
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            return [f"archive corrompue : {bad}"]
        for name in z.namelist():
            if not name.endswith((".xml", ".rels")):
                continue
            try:
                root = etree.fromstring(z.read(name))
            except etree.XMLSyntaxError as exc:
                out.append(f"{name}: XML invalide ({exc})")
                continue
            for el in root.iter():
                lname = _local(el.tag)
                for attr in _REQUIRED_ATTRS.get(lname, ()):
                    if el.get(attr) is None:
                        out.append(f"{name}: <{lname}> sans attribut requis "
                                   f"'{attr}'")
                order = _ORDERS.get(lname)
                if not order:
                    continue
                seen = -1
                for child in el:
                    cname = _local(child.tag)
                    if cname not in order:
                        continue                      # enfant hors sequence connue
                    idx = order.index(cname)
                    if idx < seen:
                        got = [_local(c.tag) for c in el]
                        out.append(f"{name}: <{lname}> ordre invalide {got} "
                                   f"(attendu la sequence {order[:4]}...)")
                        break
                    seen = idx
    return out


def _token_drift(snapshot, current, path="") -> list[str]:
    """Compare deux arbres de tokens et decrit chaque ecart, chemin compris."""
    out = []
    if isinstance(snapshot, dict) and isinstance(current, dict):
        for key in sorted(set(snapshot) | set(current)):
            where = f"{path}.{key}" if path else key
            if key not in current:
                out.append(f"{where} supprime")
            elif key not in snapshot:
                out.append(f"{where} ajoute ({current[key]!r})")
            else:
                out += _token_drift(snapshot[key], current[key], where)
    elif snapshot != current:
        out.append(f"{path} : {snapshot!r} -> {current!r}")
    return out


def _fonts_embedded(path: Path) -> bool:
    import zipfile
    with zipfile.ZipFile(path) as z:
        pres = z.read("ppt/presentation.xml").decode("utf-8", "ignore")
        has_decl = "embeddedFont" in pres
        has_part = any(n.startswith("ppt/fonts/") for n in z.namelist())
    return has_decl and has_part


# ----------------------------- CLI -----------------------------

PLANS = ["tests/charte_v2_modeles.json", "tests/rapport_hebdomadaire.json",
         "tests/ecommerce_analysis.json", "tests/diagnostic_recreation.json"]


def _image_fixture(outdir: Path) -> Path | None:
    """Fabrique des images de test deterministes + le plan qui les exerce.

    Couvre le vrai piege du flux images : une capture PORTRAIT doit etre bridee
    par la hauteur, sinon elle sort de la diapositive. Genere localement (pas de
    binaire versionne, pas de reseau).
    """
    try:
        from PIL import Image, ImageDraw
    except Exception:
        return None
    import json
    imgs = outdir / "test_images"
    imgs.mkdir(parents=True, exist_ok=True)

    def mock(name, size, mode="RGB", bg=(255, 255, 255)):
        im = Image.new(mode, size, bg if mode == "RGB" else (0, 0, 0, 0))
        d = ImageDraw.Draw(im)
        w, h = size
        d.rectangle([0, 0, w, int(h * 0.08)], fill=(25, 25, 25))
        for k in range(6):
            y = int(h * 0.14) + k * int(h * 0.11)
            d.rectangle([int(w * 0.06), y, int(w * (0.5 + 0.06 * (k % 4))), y + int(h * 0.05)],
                        fill=(231, 231, 231))
        d.rectangle([int(w * 0.06), int(h * 0.82), int(w * 0.55), int(h * 0.90)],
                    fill=(236, 67, 36))
        path = imgs / name
        im.save(path)
        return str(path)

    portrait = mock("capture_mobile.png", (750, 1600))
    paysage = mock("page_desktop.png", (1440, 900))
    carre = mock("carre.png", (900, 900))
    alpha = mock("maquette_alpha.png", (1000, 620), mode="RGBA")

    plan = {"slides": [
        {"layout": "cover", "title": "Test du rendu des images",
         "subtitle": "Capture, audit et split - portrait, paysage, carre, PNG alpha"},
        {"layout": "capture", "title": "Capture portrait 750 x 1600",
         "intro": "Points de blocage identifies", "image": portrait,
         "frictions": [{"label": "Le bouton de validation est sous la ligne de flottaison"},
                       {"label": "Les frais de port ne sont pas indiques clairement"}]},
        {"layout": "capture", "title": "Capture paysage 1440 x 900", "image": paysage,
         "frictions": [{"label": "Bandeau de tete trop haut sur mobile", "impact": "MOYEN"}]},
        {"layout": "audit", "title": "Audit avec capture portrait", "subtitle": "Page panier",
         "image": portrait, "summary": "Le tunnel perd 12 % des sessions avant paiement.",
         "impact": [{"label": "Taux actuel", "value": "2,35 %"},
                    {"label": "Cible", "value": "3,10 %"}],
         "frictions": [{"label": "Frais affiches trop tard", "impact": "ELEVE"},
                       {"label": "Champ code promo trop visible", "impact": "MOYEN"}],
         "annotations": [{"n": 1, "x": 0.5, "y": 0.35}, {"n": 2, "x": 0.6, "y": 0.7}]},
        {"layout": "split", "title": "Split - image paysage a gauche",
         "left": {"image": paysage},
         "right": {"blocks": [{"heading": "Constat", "body": "*Le visuel de tete* mange la hauteur utile."},
                              {"heading": "Recommandation", "body": "Reduire le bandeau a 40 % du viewport."}]}},
        {"layout": "split", "title": "Split - carre a droite", "swap": True, "ratio": "third",
         "left": {"image": carre},
         "right": {"cards": [{"heading": "Panier", "body": "Abandon a 68 % sur mobile."},
                             {"heading": "Paiement", "body": "Trois etapes au lieu d'une."}]}},
        {"layout": "split", "title": "Split - PNG avec transparence",
         "left": {"image": alpha},
         "right": {"blocks": [{"heading": "Maquette", "body": "*Le PNG alpha* doit rester net."}]}},
        {"layout": "capture", "title": "Capture absente : placeholder attendu",
         "frictions": [{"label": "Aucune capture fournie", "impact": "FAIBLE"}]},
        {"layout": "closing", "url": "www.the-oz.com"},
    ]}
    plan_path = outdir / "images_fixture.json"
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    return plan_path


def _render_all() -> list[Path]:
    outdir = ROOT / "out"
    outdir.mkdir(exist_ok=True)
    made = []
    py = str(ROOT / ".venv" / "bin" / "python")
    if not Path(py).exists():
        py = sys.executable
    plans = [ROOT / p for p in PLANS]
    fixture = _image_fixture(outdir)
    if fixture:
        plans.append(fixture)
    for src in plans:
        if not src.exists():
            continue
        dst = outdir / (src.stem + ".pptx")
        subprocess.run([py, str(ROOT / "engine" / "cli.py"), str(src), str(dst)],
                       check=True, capture_output=True)
        made.append(dst)
    return made


def main() -> int:
    args = sys.argv[1:]
    if not args or args[0] == "--all":
        targets = _render_all()
    else:
        targets = [Path(a) for a in args]
    reports = [check_deck(t) for t in targets]

    # Tokens de charte : l'instantane fige fait echouer toute derive silencieuse
    # d'un token de `Brand`. Une modification volontaire se valide en regenerant
    # l'instantane ET en bumpant CHARTE_VERSION.
    rep = Report(ROOT / "tests" / "charte_tokens.json")
    snap_path = ROOT / "tests" / "charte_tokens.json"
    if snap_path.exists():
        import json as _json
        snapshot = _json.loads(snap_path.read_text(encoding="utf-8"))
        current = export_tokens()
        drift = _token_drift(snapshot, current)
        rep.check("tokens de charte conformes a l'instantane fige", not drift,
                  "; ".join(drift[:4]))
    else:
        rep.check("instantane des tokens present", False,
                  "tests/charte_tokens.json absent : "
                  "`python -m oz_deck --export json > tests/charte_tokens.json`")
    reports.append(rep)

    # Les templates livres comme assets sont fabriques en XML a la main :
    # on verifie au moins leur validite structurelle.
    for asset in sorted((ROOT / "engine" / "oz_deck" / "assets").glob("oz_template*.pptx")):
        rep = Report(asset)
        defects = _ooxml_defects(asset)
        rep.check("structure OOXML valide (PowerPoint n'a rien a reparer)",
                  not defects, "; ".join(defects[:3]))
        reports.append(rep)

    for rep in reports:
        print(rep.render())
    ko = sum(len(r.failed) for r in reports)
    total = sum(len(r.rows) for r in reports)
    print(f"\n{total - ko}/{total} controles OK sur {len(reports)} deck(s).")
    if ko:
        print("REGLES VIOLEES :")
        for rep in reports:
            for rule, _, detail in rep.failed:
                print(f"  - [{rep.path.name}] {rule}" + (f" -> {detail}" if detail else ""))
    return 1 if ko else 0


if __name__ == "__main__":
    sys.exit(main())
